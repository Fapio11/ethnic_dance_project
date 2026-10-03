from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


CODE_PATH = Path(r"D:\ethnic_dance_project\page_修改版.tsx")
OUTPUT_PATH = Path(r"D:\ethnic_dance_project\page代碼_修改版.docx")


def set_run_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    fonts.set(qn("w:eastAsia"), name)


document = Document()
section = document.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
section.top_margin = Cm(1.2)
section.bottom_margin = Cm(1.2)
section.left_margin = Cm(1.35)
section.right_margin = Cm(1.35)

title = document.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(7)
title_run = title.add_run("民族舞蹈平台页面代码修改版")
set_run_font(title_run, "Microsoft YaHei", 18, bold=True)

intro = document.add_paragraph()
intro.paragraph_format.space_after = Pt(8)
intro_run = intro.add_run(
    "以下代码可整体替换 web/app/page.tsx。已修正资源核验状态字段、"
    "标准动作动态选择、reference_id 提交和多样本辅助解说，并通过 TypeScript 检查。"
)
set_run_font(intro_run, "Microsoft YaHei", 9)

code_lines = CODE_PATH.read_text(encoding="utf-8").splitlines()

for line in code_lines:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.keep_together = False
    paragraph.paragraph_format.keep_with_next = False
    code_run = paragraph.add_run(line if line else " ")
    set_run_font(code_run, "Consolas", 7.5)

document.save(OUTPUT_PATH)
print(OUTPUT_PATH)
